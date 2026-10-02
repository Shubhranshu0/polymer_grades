#!/usr/bin/env python3
"""Inspect the POLYMER COMPETITION xlsx file."""
import openpyxl

wb = openpyxl.load_workbook("/home/z/my-project/upload/POLYMER COMPETITION.xlsx", data_only=True)
print("Sheets:", wb.sheetnames)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== Sheet: {sheet_name} | dims: {ws.dimensions} | rows: {ws.max_row} cols: {ws.max_column} ===")
    # print first 30 rows
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= 30:
            print("... (truncated)")
            break
        vals = [str(v)[:30] if v is not None else "" for v in row]
        if any(vals):
            print(f"R{i+1}: {' | '.join(vals)}")
