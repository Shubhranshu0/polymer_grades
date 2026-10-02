#!/usr/bin/env python3
"""Extract POLYMER COMPETITION.xlsx -> competition.json"""
import openpyxl, json, re

SRC = "/home/z/my-project/upload/POLYMER COMPETITION.xlsx"
OUT = "/home/z/my-project/extracted/competition.json"

def clean(s):
    if s is None:
        return ""
    s = str(s).replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip()

wb = openpyxl.load_workbook(SRC, data_only=True)
print("Sheets:", wb.sheetnames)

data = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = []
    header = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = [clean(v) for v in row]
        if not any(vals):
            continue
        if i == 0:
            header = vals
            continue
        if all(not v for v in vals):
            continue
        # company, polymer, grade columns
        if not vals[0]:
            continue
        rows.append(vals)
    data[sheet_name] = {
        "header": [h for h in header if h] if header else [],
        "rows": rows,
    }
    print(f"{sheet_name}: {len(rows)} rows, header={len([h for h in header if h])} cols")

with open(OUT, "w") as f:
    json.dump(data, f, indent=1)

# summary
for sn, d in data.items():
    companies = {}
    for r in d["rows"]:
        comp = r[0]
        companies[comp] = companies.get(comp, 0) + 1
    print(sn, "->", companies)
